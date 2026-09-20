"""Real local files verify Office readers without provider accounts or an LLM."""
from pathlib import Path

from pipeline import ingest
from feeders.gdrive import sync as drive
from feeders.upload import sync as upload


def _docx(path):
    from docx import Document
    document = Document()
    document.add_heading("Project Cedar", 0)
    document.add_paragraph("Decision: launch in October.")
    cells = document.add_table(rows=1, cols=2).rows[0].cells
    cells[0].text, cells[1].text = "Owner", "Asha"
    document.save(path)


def _xlsx(path):
    from openpyxl import Workbook
    workbook = Workbook()
    workbook.active.title = "Budget"
    workbook.active.append(["Item", "Amount"])
    workbook.active.append(["Pilot", 1200])
    workbook.active.append(["Total", "=SUM(B2:B2)"])
    workbook.create_sheet("Decisions").append(["Launch", "October"])
    workbook.save(path)


def _pptx(path):
    from pptx import Presentation
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Project Cedar"
    slide.placeholders[1].text = "Launch in October"
    slide.notes_slide.notes_text_frame.text = "Asha owns the pilot."
    deck.save(path)


def _pdf(path):
    # A minimal real PDF avoids introducing a fixture-generation dependency.
    from pypdf import PdfWriter
    from pypdf.generic import (NameObject, DictionaryObject, DecodedStreamObject)
    writer = PdfWriter()
    page = writer.add_blank_page(width=400, height=200)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                             NameObject("/Subtype"): NameObject("/Type1"),
                             NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"):
        DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
    content = DecodedStreamObject()
    content.set_data(b"BT /F1 12 Tf 20 100 Td (Project Cedar launches in October.) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(content)
    writer.write(path)


def test_native_docx_preserves_paragraphs_and_tables(tmp_path, monkeypatch):
    path = tmp_path / "decision.docx"
    _docx(path)
    monkeypatch.setattr(ingest, "have", lambda _: False)
    text, method = ingest.extract_binary(path)
    assert method == "python-docx"
    assert "Project Cedar" in text and "launch in October" in text
    assert "Owner\tAsha" in text


def test_xlsx_preserves_every_sheet_and_labels_unevaluated_formulas(tmp_path):
    path = tmp_path / "budget.xlsx"
    _xlsx(path)
    text, method = ingest.extract_binary(path)
    assert method == "openpyxl"
    assert "Sheet: Budget" in text and "Pilot\t1200" in text
    assert "Sheet: Decisions" in text and "Launch\tOctober" in text
    assert "[formula: =SUM(B2:B2)]" in text


def test_pptx_includes_slide_text_and_speaker_notes(tmp_path):
    path = tmp_path / "plan.pptx"
    _pptx(path)
    text, method = ingest.extract_binary(path)
    assert method == "python-pptx"
    assert "Slide 1" in text and "Launch in October" in text
    assert "Speaker notes:\nAsha owns the pilot." in text


def test_actual_pdf_extracts_with_installed_runtime_and_python_fallback(tmp_path, monkeypatch):
    path = tmp_path / "decision.pdf"
    _pdf(path)
    text, method = ingest.extract_binary(path)
    assert "Project Cedar launches in October." in text
    assert method in {"pdftotext", "pypdf"}
    monkeypatch.setattr(ingest, "have", lambda _: False)
    text, method = ingest.extract_binary(path)
    assert method == "pypdf" and "Project Cedar launches in October." in text


def test_upload_and_drive_extract_actual_office_bytes(tmp_path, monkeypatch):
    for extension, make_file, expected in [
        ("docx", _docx, "Project Cedar"),
        ("xlsx", _xlsx, "Pilot\t1200"),
        ("pptx", _pptx, "Launch in October"),
    ]:
        path = tmp_path / f"fixture.{extension}"
        make_file(path)
        kind, text = upload._dispatch_extract(path)
        assert kind == "binary_doc" and expected in text
        mime = next(mime for mime, suffix in drive._OFFICE_MIMES.items() if suffix == path.suffix)
        monkeypatch.setattr(drive, "_get_bytes", lambda _, path=path: path.read_bytes())
        document = {"name": path.name, "id": "fixture", "mime_type": mime}
        assert expected in drive.export_text(document)
        assert drive._source_type(document) == "binary_doc"
        assert drive._wanted({"name": path.name, "mimeType": mime})
