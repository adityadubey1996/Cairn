#!/usr/bin/env python3
"""
ingest.py — the `/wiki ingest` step for a codebase.

The wiki skill expects `raw/entries/{date}_{id}.md`, one markdown file per logical
entry, with YAML frontmatter carrying a date. A codebase has no entries and no
dates, so this script manufactures both:

  entry  = a *source unit*, not a file. 2,479 tracked paths collapse to ~250 units.
  date   = the commit date the unit first appeared. Absorbing chronologically then
           replays how the system was built, which is the only ingest order that
           makes rival implementations legible instead of contradictory.

Idempotent: same repo state in, same files out.

    pip install -r requirements.txt   # Python PDF and Office readers
    brew install pandoc poppler       # optional preferred DOCX/PDF readers
    python .cursor/skills/wiki/scripts/ingest.py --repo . --out raw/entries

Flags:
    --code-depth N      max path depth for a code package unit (default 3)
    --max-package-files N   split a package deeper while it exceeds N files (default 120)
    --bulk-threshold N  N+ data files in one dir become a single dataset unit (default 5)
    --check             don't write; report which units' SHAs moved since last run

Project excludes: repo-root .wikiignore (gitignore-lite). See that file for
backend/, Asset-Manager/, etc. Hardcoded EXCLUDE_DIR/EXCLUDE_EXT still apply.
"""

import argparse, hashlib, json, os, re, shutil, subprocess, sys
from collections import defaultdict
from datetime import datetime
from fnmatch import fnmatch
from pathlib import Path

# ------------------------------------------------------------------ what counts

EXCLUDE_DIR = [
    "graphify-out",                  # stale AST snapshot of a prior tree; excluded until regenerated
    "node_modules", ".venv", "__pycache__", "dist", "build", ".git",
    "migrations/versions",
]
EXCLUDE_EXT = {".png", ".jpg", ".jpeg", ".gif", ".ico", ".svg", ".wasm", ".gpkg",
               ".mbtiles", ".lock", ".pyc", ".map", ".woff", ".woff2"}

PROSE_EXT   = {".md", ".rst", ".txt"}
BINARY_EXT  = {".docx", ".pptx", ".pdf", ".xlsx"}
DIAGRAM_EXT = {".excalidraw", ".mmd"}
CODE_EXT    = {".py", ".ts", ".tsx", ".js", ".jsx", ".sol", ".sql", ".daml"}
DATA_EXT    = {".json", ".csv", ".tsv", ".parquet", ".geojson"}

# What actually runs: ports, processes, env, proxies. A `runtime_contract` unit
# becomes an article, so the bar is "reading this tells you how the thing boots".
#
# package.json is deliberately NOT here. It is a dependency manifest, and asking
# for an article about one buys npm metadata: six were generated saying things
# like "the package has a version of 0.0.0 and is set to private". Dropping it
# here sends it to the `data` kind — still citable by a real runtime article,
# never a subject of its own. It stays in PKG_MARKER below, which is a different
# job: finding package boundaries in JS repos.
CONTRACT_RE = re.compile(
    r"(docker-compose[\w.-]*\.ya?ml|\.github/workflows/.+|Dockerfile[\w.-]*|"
    r"pyproject\.toml|.*nginx.*\.conf|martin[\w.-]*\.ya?ml|"
    r"alembic\.ini|\.env\.example|Makefile)$"
)
PKG_MARKER = {"__init__.py", "pyproject.toml", "package.json", "setup.py"}

# A code_package sha is agg(git blob shas of its files) — the generated fact
# sheet is not part of it, so changing the fact-sheet format alone leaves the
# absorb queue empty. Bump this to push every code package into `changed` and
# let the existing drift machinery schedule the re-absorb.
FACTSHEET_VERSION = "2-graph-edges"
SUPERSEDED_RE = re.compile(r"(/BackUps?/|/archive/|/old/|\bcopy\b|_Older_Version|_deprecated|_bak)", re.I)
VERSION_RE = re.compile(r"[_-]?v(\d+)(?=\.|$)", re.I)

WIKIIGNORE_NAME = ".wikiignore"


def load_wikiignore(repo):
    """Load repo-root .wikiignore patterns (gitignore-lite). Missing file → []."""
    path = Path(repo) / WIKIIGNORE_NAME
    if not path.is_file():
        return []
    patterns = []
    for line in path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#"):
            continue
        patterns.append(line.replace("\\", "/"))
    return patterns


def wikiignored(rel, patterns):
    """True if repo-relative path matches a .wikiignore pattern."""
    if not patterns:
        return False
    p = rel.replace("\\", "/")
    name = Path(p).name
    for pat in patterns:
        if pat.endswith("/"):
            root = pat.rstrip("/")
            if p == root or p.startswith(root + "/"):
                return True
            continue
        if any(c in pat for c in "*?[]"):
            if fnmatch(p, pat) or fnmatch(name, pat):
                return True
            continue
        root = pat.rstrip("/")
        if p == root or p.startswith(root + "/"):
            return True
    return False


def classify(rel, wikiignore=None):
    p = rel.replace("\\", "/")
    if wikiignored(p, wikiignore or []):
        return None
    if any(p.startswith(d) or f"/{d}/" in f"/{p}" for d in EXCLUDE_DIR):
        return None
    ext = Path(p).suffix.lower()
    if ext in EXCLUDE_EXT:
        return None
    if p.endswith(".excalidraw copy") or ext in DIAGRAM_EXT:
        return "diagram"
    if ext in BINARY_EXT:
        return "binary_doc"
    if CONTRACT_RE.search(p):
        return "runtime_contract"
    if ext in PROSE_EXT:
        return "doc"
    if not ext and Path(p).name.lower() in {'readme', 'license', 'licence', 'changelog', 'authors', 'contributing'}:
        return 'doc'
    if ext in CODE_EXT:
        return "code"
    if ext in DATA_EXT or ext in {".yml", ".yaml", ".toml", ".ini", ".conf"}:
        return "data"
    return None


# ------------------------------------------------------------------ git

def git(repo, *args, **kw):
    return subprocess.run(["git", "-C", str(repo), *args],
                          check=True, capture_output=True, text=True,
                          errors="replace", **kw).stdout


def tracked(repo):
    out = []
    for line in git(repo, "ls-files", "-s").splitlines():
        meta, path = line.split("\t", 1)
        out.append((meta.split()[1], path))
    return out


def history(repo):
    """
    One pass over the log to get, per path: first-seen date, last-touched date,
    commit count, authors. Cheaper and more reliable than --follow per file.

    -M with --name-status (not --name-only) is what makes a rename visible as
    an R### line instead of a plain delete-and-add — without it, a renamed
    file's "first-seen" date reset to the rename commit instead of its true
    origin, which corrupts the chronological absorb order this function
    exists to support.
    """
    first, last, count, authors = {}, {}, defaultdict(int), defaultdict(set)
    raw = git(repo, "log", "--reverse", "-M", "--name-status",
              "--format=@@%H|%aI|%an", "--no-merges")
    date = author = None
    for line in raw.splitlines():
        if line.startswith("@@"):
            _, date, author = line[2:].split("|", 2)
            continue
        if not line.strip() or date is None:
            continue
        parts = line.rstrip("\n").split("\t")
        status = parts[0]
        if status.startswith("R") and len(parts) == 3:
            old_p, p = parts[1], parts[2]
            first[p] = first.get(old_p, date)
            count[p] += count.get(old_p, 0)
            authors[p] |= authors.get(old_p, set())
        else:
            p = parts[-1]
            first.setdefault(p, date)
        last[p] = date
        count[p] += 1
        authors[p].add(author)
    return first, last, count, authors


# ------------------------------------------------------------------ extractors

_HAVE_CACHE = {}


def have(cmd):
    # shutil.which() alone is fooled by pyenv shims: the shim file exists on PATH
    # but resolves to a different env and exits 127 ("command not found"), so
    # extract_binary() picked markitdown and silently produced EXTRACTION FAILED.
    # Run the tool once and cache the verdict.
    if cmd not in _HAVE_CACHE:
        ok = shutil.which(cmd) is not None
        if ok:
            try:
                ok = subprocess.run(
                    [cmd, "--version"], capture_output=True, timeout=15
                ).returncode != 127
            except (OSError, subprocess.SubprocessError):
                ok = False
        _HAVE_CACHE[cmd] = ok
    return _HAVE_CACHE[cmd]


def _extract_office(path: Path):
    """Text, tables, and spreadsheet formulas; never execute document content."""
    if path.suffix.lower() == ".docx":
        from docx import Document
        from docx.table import Table
        chunks = []
        for block in Document(path).iter_inner_content():
            if isinstance(block, Table):
                chunks.append("\n".join("\t".join(cell.text for cell in row.cells)
                                         for row in block.rows))
            else:
                chunks.append(block.text)
        return "\n\n".join(chunks), "python-docx"
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=False)
        chunks = []
        has_values = False
        try:
            for sheet in workbook:
                chunks.append(f"## Sheet: {sheet.title}")
                for row in sheet.iter_rows():
                    values = [(f"[formula: {cell.value}]" if cell.data_type == "f"
                               else str(cell.value) if cell.value is not None else "")
                              for cell in row]
                    if any(values):
                        has_values = True
                        chunks.append("\t".join(values).rstrip("\t"))
        finally:
            workbook.close()
        return "\n".join(chunks) if has_values else "", "openpyxl"
    if path.suffix.lower() == ".pptx":
        from pptx import Presentation

        def texts(shapes):
            for shape in shapes:
                if shape.has_text_frame:
                    yield shape.text
                if shape.has_table:
                    yield "\n".join("\t".join(cell.text for cell in row.cells)
                                      for row in shape.table.rows)
                if hasattr(shape, "shapes"):
                    yield from texts(shape.shapes)

        chunks = []
        has_text = False
        for number, slide in enumerate(Presentation(path).slides, 1):
            chunks.append(f"## Slide {number}")
            slide_text = [text for text in texts(slide.shapes) if text.strip()]
            has_text = has_text or bool(slide_text)
            chunks.extend(slide_text)
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame
                if notes is not None and notes.text.strip():
                    has_text = True
                    chunks.append(f"Speaker notes:\n{notes.text}")
        return "\n\n".join(chunks) if has_text else "", "python-pptx"
    raise ValueError(f"unsupported Office format: {path.suffix}")


def extract_binary(path: Path):
    """Text from docx/pptx/pdf/xlsx, including native Python fallbacks."""
    path = Path(path)
    ext = path.suffix.lower()
    errors = []
    preferred = {
        ".docx": ("pandoc", ["pandoc", "--standalone", "-t", "markdown", str(path)]),
        ".pdf": ("pdftotext", ["pdftotext", "-layout", str(path), "-"]),
    }.get(ext)
    if preferred and have(preferred[0]):
        try:
            text = subprocess.run(preferred[1], capture_output=True, text=True,
                                  check=True, timeout=60).stdout
            if text.strip():
                return text, preferred[0]
        except Exception as error:
            errors.append(f"{preferred[0]}: {error}")
    try:
        if ext in {".docx", ".pptx", ".xlsx"}:
            return _extract_office(path)
        if ext == ".pdf":
            from pypdf import PdfReader
            text = "\n\n".join((page.extract_text() or "") for page in PdfReader(str(path)).pages)
            return text, "pypdf" if text.strip() else "NO_TEXT_LAYER: PDF needs OCR"
    except Exception as error:
        errors.append(str(error))
    if have("markitdown"):
        try:
            return subprocess.run(["markitdown", str(path)], capture_output=True,
                                  text=True, check=True, timeout=60).stdout, "markitdown"
        except Exception as error:
            errors.append(f"markitdown: {error}")
    return "", f"FAILED: {'; '.join(errors)}" if errors else "NO_EXTRACTOR"


def extract_excalidraw(path: Path):
    """
    Excalidraw files are JSON. Pulling the text elements out turns 25 opaque
    diagrams into readable flow descriptions, ordered top-to-bottom then
    left-to-right so the reading order roughly matches the drawn flow.
    """
    try:
        doc = json.loads(path.read_text(encoding="utf-8", errors="replace"))
    except Exception as e:
        return f"_could not parse: {e}_"
    els = [e for e in doc.get("elements", []) if e.get("type") == "text" and e.get("text")]
    els.sort(key=lambda e: (round(e.get("y", 0) / 60), e.get("x", 0)))
    arrows = sum(1 for e in doc.get("elements", []) if e.get("type") in ("arrow", "line"))
    body = "\n".join(f"- {e['text'].strip()}" for e in els if e["text"].strip())
    return f"_{len(els)} labels, {arrows} connectors._\n\n{body}" if body else "_no text labels_"


def summarize_dataset(repo, files):
    lines = [f"_{len(files)} files._", ""]
    for rel in files[:4]:
        p = repo / rel
        try:
            if p.suffix.lower() in {".csv", ".tsv"}:
                head = p.read_text(errors="replace").splitlines()[:2]
                lines += [f"**{rel}**", "```", *head, "```", ""]
            elif p.suffix.lower() == ".json":
                d = json.loads(p.read_text(errors="replace"))
                keys = list(d)[:25] if isinstance(d, dict) else f"array[{len(d)}]"
                lines += [f"**{rel}** — top-level keys: `{keys}`", ""]
        except Exception:
            lines.append(f"**{rel}** — unreadable")
    if len(files) > 4:
        lines.append(f"_… and {len(files) - 4} more._")
    return "\n".join(lines)


IMPORT_RE = re.compile(r"^\s*(?:from\s+([.\w]+)|import\s+([.\w]+))", re.M)
DEF_RE = re.compile(r"^(?:class|def|async def|export (?:default )?(?:function|const|class))\s+(\w+)", re.M)
ROUTE_RE = re.compile(r"@\w*\.?(get|post|put|patch|delete)\(\s*[\"']([^\"']+)", re.I)


def graph_edges_block(pkg, edges):
    """Cross-package structure for `pkg`, as an itinerary for the absorb step.

    The import histogram above keys on `mod.split(".")[0]`, so it cannot say
    which package a symbol came from and never says who calls in. Without that
    the writer sees a package in isolation, cannot trace a flow, and paraphrases
    whatever doc narrates it — module articles measured 2% code-grounded.

    These lines are navigation, not content. The article must not assert
    "called by workflows" (SKILL.md rule 11: recomputable → query the graph).
    """
    e = (edges or {}).get(pkg)
    if not e:
        return []
    out = []
    if e["inbound"]:
        out.append("**Called by:** " + ", ".join(
            f"`{k}`({v})" for k, v in e["inbound"].most_common(6)))
    if e["outbound"]:
        out.append("**Calls into:** " + ", ".join(
            f"`{k}`({v})" for k, v in e["outbound"].most_common(6)))
    if e["entry_points"]:
        out += ["", "**Open next** — what the rest of the system actually calls:"]
        for label, loc, n, callers in e["entry_points"]:
            via = f" via {', '.join(callers)}" if callers else ""
            out.append(f"- `{label}` — `{loc}` ← {n} external{via}")
    return out + [""] if out else []


def fact_sheet(repo, pkg, files, edges=None):
    """
    A code package is never pasted into an entry. It gets a fact sheet: enough
    for the absorb step to know what the package is and what to open next.
    """
    loc = 0; defs = []; routes = []; imports = defaultdict(int); tests = []
    for rel in files:
        p = repo / rel
        try:
            src = p.read_text(errors="replace")
        except Exception:
            continue
        loc += src.count("\n")
        if "test" in Path(rel).name.lower():
            tests.append(rel)
        defs += [(rel, n) for n in DEF_RE.findall(src)[:8]]
        routes += [f"{m.upper()} {u}" for m, u in ROUTE_RE.findall(src)]
        for a, b in IMPORT_RE.findall(src):
            mod = (a or b).split(".")[0].lstrip(".")
            if mod:
                imports[mod] += 1
    ext_c = defaultdict(int)
    for rel in files:
        ext_c[Path(rel).suffix] += 1

    out = [f"**Package** `{pkg}` — {len(files)} files, ~{loc} lines, "
           f"{len(tests)} test files.", "",
           "**Composition:** " + ", ".join(f"{v}× `{k}`" for k, v in sorted(ext_c.items(), key=lambda x: -x[1])), ""]
    readme = next((f for f in files if Path(f).name.lower().startswith("readme")), None)
    if readme:
        out += ["**Has README:** `%s` — read it before writing the article." % readme, ""]
    if routes:
        out += ["**HTTP surface:**"] + [f"- `{r}`" for r in sorted(set(routes))[:25]] + [""]
    if imports:
        top = sorted(imports.items(), key=lambda x: -x[1])[:15]
        out += ["**Imports (by frequency):** " + ", ".join(f"`{k}`×{v}" for k, v in top), ""]
    if defs:
        out += ["**Declared symbols (sample):**"]
        for rel, n in defs[:30]:
            out.append(f"- `{n}` — `{rel}`")
        out.append("")
    out += graph_edges_block(pkg, edges)
    out += ["**Files:**"] + [f"- `{f}`" for f in files[:60]]
    if len(files) > 60:
        out.append(f"- _… and {len(files)-60} more_")
    return "\n".join(out)


# ------------------------------------------------------------------ units

def pkg_root(repo, rel, cap):
    parts = Path(rel).parts[:-1]
    capped = str(Path(*parts[:cap])) if parts else "."
    cur = Path(rel).parent
    while str(cur) not in (".", "/", ""):
        if any((repo / cur / m).exists() for m in PKG_MARKER):
            return capped if len(cur.parts) > cap else str(cur)
        cur = cur.parent
    return capped


def split_oversized(packages, limit):
    """pkg_root stops at the nearest PKG_MARKER, which assumes markers are dense.
    They are in Python (__init__.py per level) and sparse in pnpm (one package.json
    per workspace package), so a 482-file SPA lands in one unit. Subdivide by path
    level until each package is small enough to write an article about.
    """
    out = {}
    for pkg, mem in packages.items():
        if len(mem) <= limit:
            out[pkg] = mem
            continue
        depth = len(Path(pkg).parts)
        deeper = defaultdict(list)
        for sha, rel in mem:
            parts = Path(rel).parts
            child = str(Path(*parts[:depth + 1])) if len(parts) > depth + 1 else pkg
            deeper[child].append((sha, rel))
        # {pkg: mem} means no deeper level exists — accepting it would loop forever.
        # A single deeper key is still progress: the package gets a more specific path.
        out.update(deeper if set(deeper) != {pkg} else {pkg: mem})
    return out


def slug(s):
    return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")[:80] or "root"


def agg(shas):
    return hashlib.sha1("".join(sorted(shas)).encode()).hexdigest()


def build_units(repo, cap, bulk, wikiignore=None, max_pkg=120):
    if wikiignore is None:
        wikiignore = load_wikiignore(repo)
    probe = subprocess.run(["git", "-C", str(repo), "rev-parse", "--is-inside-work-tree"],
                           capture_output=True, text=True)
    if probe.returncode != 0:
        return []
    first, last, ncommit, authors = history(repo)
    rows = [(sha, rel, k) for sha, rel in tracked(repo)
            if (k := classify(rel, wikiignore)) is not None]

    dircount = defaultdict(int)
    for _s, rel, k in rows:
        if k == "data":
            dircount[str(Path(rel).parent)] += 1
    bulk_dirs = {d for d, n in dircount.items() if n >= bulk}

    units, packages, datasets = [], defaultdict(list), defaultdict(list)

    def dates(paths):
        f = min((first.get(p, "") for p in paths if first.get(p)), default="")
        l = max((last.get(p, "") for p in paths if last.get(p)), default="")
        return f, l

    for sha, rel, kind in rows:
        if kind == "code":
            packages[pkg_root(repo, rel, cap)].append((sha, rel)); continue
        if kind == "data" and str(Path(rel).parent) in bulk_dirs:
            datasets[str(Path(rel).parent)].append((sha, rel)); continue
        f, l = dates([rel])
        units.append(dict(id=slug(rel), kind=kind, path=rel, sha=sha, files=[rel],
                          first=f, last=l, commits=ncommit.get(rel, 0),
                          authors=sorted(authors.get(rel, [])),
                          status="superseded" if SUPERSEDED_RE.search("/" + rel) else "active"))

    while True:
        split = split_oversized(packages, max_pkg)
        if split.keys() == packages.keys():
            break
        packages = split

    for pkg, mem in packages.items():
        paths = [r for _s, r in mem]; f, l = dates(paths)
        units.append(dict(id="pkg-" + slug(pkg), kind="code_package", path=pkg,
                          sha=agg([s for s, _ in mem] + [FACTSHEET_VERSION]),
                          files=sorted(paths), first=f, last=l,
                          commits=sum(ncommit.get(p, 0) for p in paths),
                          authors=sorted({a for p in paths for a in authors.get(p, set())}),
                          status="active"))
    for d, mem in datasets.items():
        paths = [r for _s, r in mem]; f, l = dates(paths)
        units.append(dict(id="data-" + slug(d), kind="dataset", path=d,
                          sha=agg(s for s, _ in mem), files=sorted(paths), first=f, last=l,
                          commits=sum(ncommit.get(p, 0) for p in paths),
                          authors=sorted({a for p in paths for a in authors.get(p, set())}),
                          status="active"))

    # version chains: API_v1/v2/v3, deck_v2/v3, and the unversioned base as v0
    fams, bases = defaultdict(list), {}
    for u in units:
        stem = Path(u["path"]).stem
        m = VERSION_RE.search(stem)
        if m:
            fams[str(Path(u["path"]).parent / VERSION_RE.sub("", stem))].append((int(m.group(1)), u))
        elif Path(u["path"]).name:
            bases[str(Path(u["path"]).with_suffix(""))] = u
    for fam, items in fams.items():
        if fam in bases:
            items.append((0, bases[fam]))
        items.sort(key=lambda x: x[0])
        newest = items[-1][1]
        for _v, u in items[:-1]:
            u["status"] = "superseded"; u["superseded_by"] = newest["id"]
        newest["supersedes"] = [u["id"] for _v, u in items[:-1]]

    units.sort(key=lambda u: (u["first"] or "9999", u["id"]))
    return units


# ------------------------------------------------------------------ rendering

def load_package_edges(repo):
    """Cross-package edges from graphify, or {} if unavailable.

    The graph is written into the clone by the graph step, so a repo ingested
    before that runs — or one where graphify failed — has none. Missing, stale
    or unreadable graph must degrade to today's behaviour — a fact sheet without
    an itinerary, never a failed ingest.
    """
    gp = repo / "graphify-out" / "graph.json"
    if not gp.is_file():
        return {}
    try:
        sys.path.insert(0, str(Path(__file__).parent))
        from graph_verify import package_edges
        return package_edges(json.loads(gp.read_text(encoding="utf-8")))
    except Exception as e:
        print(f"  graph: skipped ({type(e).__name__}: {e}) — fact sheets unenriched")
        return {}


def body_for(repo, u, edges=None):
    p = repo / u["path"]
    if u["kind"] in ("doc", "runtime_contract"):
        try:
            return p.read_text(errors="replace")
        except Exception as e:
            return f"_unreadable: {e}_"
    if u["kind"] == "diagram":
        return extract_excalidraw(p)
    if u["kind"] == "binary_doc":
        text, method = extract_binary(p)
        if not text.strip():
            return (f"> **EXTRACTION FAILED** ({method}). Do not absorb this entry — "
                    f"install pandoc/markitdown and re-run ingest, or the wiki will "
                    f"silently lose `{u['path']}`.")
        return f"<!-- extracted via {method} -->\n\n{text}"
    if u["kind"] == "code_package":
        return fact_sheet(repo, u["path"], u["files"], edges)
    if u["kind"] == "dataset":
        return summarize_dataset(repo, u["files"])
    return "_no body_"


def yaml_list(xs):
    return "[" + ", ".join(json.dumps(x) for x in xs) + "]"


def render(repo, u, edges=None, head=None):
    date = (u["first"] or "1970-01-01T00:00:00")[:10]
    time = (u["first"] or "T00:00:00").split("T")[1][:8]
    fm = [
        "---",
        f"id: {u['id']}",
        f"date: {date}",
        f'time: "{time}"',
        f"source_type: {u['kind']}",
        f"path: {json.dumps(u['path'])}",
        f"sha: {u['sha']}",
        f"last_touched: {(u['last'] or '')[:10]}",
        f"commits: {u['commits']}",
        f"authors: {yaml_list(u['authors'][:8])}",
        f"status: {u['status']}",
    ]
    if u.get("superseded_by"):
        fm.append(f"superseded_by: {u['superseded_by']}")
    if u.get("supersedes"):
        fm.append(f"supersedes: {yaml_list(u['supersedes'])}")
    if u["kind"] in ("code_package", "dataset"):
        fm.append(f"file_count: {len(u['files'])}")
    # Which revision this entry describes. Articles absorbed from it inherit the
    # value as `built_from_commit`, so a reader can tell what tree a claim was
    # written against without reconstructing it from scattered source shas.
    if head:
        fm.append(f"built_from_commit: {head}")
    fm += ["tags: []", "---", ""]
    header = f"# `{u['path']}`\n"
    if u["status"] == "superseded":
        header += ("\n> **SUPERSEDED**"
                   + (f" by `{u['superseded_by']}`." if u.get("superseded_by") else ".")
                   + " Absorb for history only. Never cite as current behaviour.\n")
    return "\n".join(fm) + header + "\n" + body_for(repo, u, edges) + "\n"


# ------------------------------------------------------------------ feeder inbox

INBOX_DIRNAME = "inbox"


INBOX_EXCLUDE_NAME = ".inbox-exclude"


def load_inbox_exclude(repo):
    """Repo-root .inbox-exclude: exact `path:` values to keep out of absorb —
    e.g. a chat day whose content shouldn't reach an LLM. Deliberately a
    separate list from .wikiignore (exact-match only, no directory/glob
    semantics): .wikiignore's bare `sources/` entry — there to stop the repo
    walk double-scanning fetched content — would otherwise match every
    inbox entry's path (all start with sources/) and silently exclude
    everything. Missing file → empty set."""
    path = Path(repo) / INBOX_EXCLUDE_NAME
    if not path.is_file():
        return set()
    return {line.strip() for line in path.read_text(encoding="utf-8").splitlines()
            if line.strip() and not line.strip().startswith("#")}


def load_inbox(inbox_dir, exclude=None):
    """Adopt pre-rendered entries written by feeders (chat, meetings, videos).

    git cannot see these sources, so build_units() can't manufacture them — and
    only the feeder knows the true date of a Google Chat thread or a recorded
    meeting. The feeder therefore owns the frontmatter; we parse just enough to
    put the unit in the manifest and the absorb queue, and copy the body through
    verbatim rather than re-rendering it.

    `raw/` is gitignored, so nothing here ever enters git history — only the
    article absorb writes from it does.
    """
    units = []
    if not inbox_dir.is_dir():
        return units
    for f in sorted(inbox_dir.glob("*.md")):
        text = f.read_text(encoding="utf-8", errors="replace")
        end = text.find("\n---", 3) if text.startswith("---") else -1
        if end == -1:
            print(f"  inbox: skipped {f.name} — missing or unterminated frontmatter")
            continue
        fm = {}
        for line in text[3:end].splitlines():
            key, sep, val = line.partition(":")
            if sep:
                fm[key.strip()] = val.strip().strip('"')
        if exclude and fm.get("path", "") in exclude:
            print(f"  inbox: skipped {f.name} — path excluded via {INBOX_EXCLUDE_NAME}")
            continue
        date = fm.get("date") or "1970-01-01"
        units.append({
            "id": fm.get("id") or slug(f.stem),
            "path": fm.get("path") or f"inbox://{f.name}",
            # feeder-supplied sha lets a re-collected thread with new messages
            # land in `changed`; fall back to hashing the entry itself
            "sha": fm.get("sha") or hashlib.sha1(text.encode()).hexdigest(),
            "kind": fm.get("source_type", "doc"),
            "status": fm.get("status", "active"),
            "first": f"{date}T{fm.get('time', '00:00:00')}",
            "last": date,
            "commits": 0,
            "authors": [x.strip().strip('"') for x in fm.get("authors", "").strip("[]").split(",")
                        if x.strip()],
            "files": [],
            "_inbox_file": f,
        })
    return units


# ------------------------------------------------------------------ main

def pair_renames(prev, cur, prev_path):
    """Match units that moved: identical content sha, different id.

    A unit id is derived from its path, so `docs/X.docx` is a different unit from
    `X.docx` and a move surfaces as removed+new. Absorbing that naively deletes the
    article and writes a second one about the same document. Pairing by sha is exact
    rather than heuristic — a move keeps the bytes identical. A file that moved *and*
    changed will not pair, and correctly stays removed+new: that really is a delete
    plus a create.

    Returns (renamed, paired_old_ids, paired_new_ids).
    """
    gone = prev.keys() - cur.keys()
    fresh = cur.keys() - prev.keys()
    by_sha = defaultdict(list)
    for i in sorted(fresh):
        by_sha[cur[i][0]].append(i)
    renamed = []
    for old in sorted(gone):
        candidates = by_sha.get(prev[old][0])
        if candidates:
            # pop so two byte-identical files can't both pair to the same new unit
            renamed.append({"from": old, "to": candidates.pop(0),
                            "from_path": prev_path.get(old, ""),
                            "sha": prev[old][0]})
    return renamed, {r["from"] for r in renamed}, {r["to"] for r in renamed}


def build_pending(units, prev, prev_path, absorbed):
    """The absorb queue: what's new, changed, removed, or renamed since the
    last ingest. `units` are this run's units (current); `prev`/`prev_path`
    are the previous manifest's {id: (sha, status)} / {id: path}; `absorbed`
    is the set of unit ids already written into an article.

    The "a renamed-and-already-absorbed unit lands in 'renamed', never also
    in 'new'" guarantee this function provides is WITHIN THIS ONE CALL only.
    It does not repoint anything durable: nothing here rewrites
    `_absorb_log.json` (still keyed by the OLD id) or the `unit:` frontmatter
    already stamped into the old article. So on the NEXT ingest run, once
    `prev`/`prev_path` have caught up to include the new id, that unit no
    longer pairs as a rename — its new id simply isn't in `absorbed`, looks
    like a never-absorbed unit, and gets queued into "new" again.
    find_existing_article() in absorb_runner.py matches on `unit: <new-id>`,
    which the old article never carries, so that second absorb produces a
    genuine duplicate article rather than integrating into the renamed
    unit's existing one. Repointing the log/frontmatter on rename detection
    would close that gap but is a larger design decision, deliberately left
    for a future task.
    """
    by_id = {u["id"]: u for u in units}
    cur = {u["id"]: (u["sha"], u["status"]) for u in units}
    renamed, _paired_old, _paired_new = pair_renames(prev, cur, prev_path)
    _new = cur.keys() - prev.keys()
    _removed = prev.keys() - cur.keys()

    # A unit that pairs as a rename goes through "renamed", never "new" —
    # unconditionally, whether or not its old id was ever absorbed. Excluding
    # it only from `_new - _paired_new` and not from `unabsorbed` is how a
    # renamed-and-already-absorbed unit used to land in BOTH buckets at once.
    if isinstance(absorbed, dict):
        # A discovery manifest describes what was seen, not what was compiled.
        # Keep an unsuccessful update queued even after another discovery pass.
        unabsorbed = (cur.keys() - absorbed.keys()) - _paired_new
        unfinished = {uid for uid in cur.keys() & absorbed.keys()
                      if absorbed[uid] != cur[uid][0]}
        newly_queued = unabsorbed
    else:
        # Compatibility for callers still supplying an old ID-only set.
        unabsorbed = (cur.keys() - absorbed) - _paired_new
        unfinished = set()
        newly_queued = (_new - _paired_new) | unabsorbed

    def chrono(ids):
        """Absorb oldest-first, per SKILL.md's own contract. A unit with no
        dated `first` (shouldn't happen; every unit gets one) sorts last."""
        def key(i):
            u = by_id.get(i)
            return (u["first"] or "9999", i) if u else ("9999", i)
        return sorted(ids, key=key)

    return {
        "new": chrono(newly_queued),
        # status counts as a change: a doc demoted to superseded needs re-absorbing
        # even though its bytes never moved
        "changed": chrono(unfinished | {k for k in cur.keys() & prev.keys()
                          if cur[k] != prev[k] and k not in unabsorbed}),
        "removed": sorted(_removed - _paired_old),
        "renamed": renamed,
        "first_run": not prev,
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--repo", type=Path, default=Path("."))
    ap.add_argument("--out", type=Path, default=Path("raw/entries"))
    ap.add_argument("--wiki", type=Path, default=None,
                    help="durable wiki dir; holds the manifest twin and absorb log")
    ap.add_argument("--source-only", action="store_true",
                    help="ingest connector inbox entries without scanning Git files")
    ap.add_argument("--code-depth", type=int, default=3)
    ap.add_argument("--bulk-threshold", type=int, default=5)
    ap.add_argument("--max-package-files", type=int, default=120)
    ap.add_argument("--check", action="store_true")
    a = ap.parse_args()

    repo = a.repo.resolve()
    wikiignore = load_wikiignore(repo)
    if not a.out.is_absolute():
        a.out = repo / a.out
    units = ([] if a.source_only else
             build_units(repo, a.code_depth, a.bulk_threshold, wikiignore=wikiignore,
                         max_pkg=a.max_package_files))
    units += load_inbox(a.out.parent / INBOX_DIRNAME, load_inbox_exclude(repo))
    manifest = a.out.parent / "_manifest.json"
    # Committed twin of the manifest. raw/ is gitignored, so on a fresh clone
    # (CI) this copy is the only baseline that exists to diff against.
    # The durable wiki may live outside the clone (ai-brain keeps it there so
    # `git reset --hard` cannot destroy paid articles), so the manifest twin
    # and the absorb log must be looked up there too.
    wiki_dir = a.wiki.resolve() if a.wiki else repo / "wiki"
    committed = wiki_dir / "_manifest.json"

    if a.check:
        base = committed if committed.exists() else manifest
        if not base.exists():
            sys.exit("no manifest yet — run ingest first")
        _m = json.loads(base.read_text())
        old = {u["id"]: (u["sha"], u.get("status", "active")) for u in _m}
        old_path = {u["id"]: u.get("path", "") for u in _m}
        new = {u["id"]: (u["sha"], u["status"]) for u in units}
        new_path = {u["id"]: u["path"] for u in units}
        renamed, paired_old, paired_new = pair_renames(old, new, old_path)
        moved = sorted(k for k in old.keys() & new.keys() if old[k] != new[k])
        added = sorted((new.keys() - old.keys()) - paired_new)
        gone = sorted((old.keys() - new.keys()) - paired_old)
        print(f"changed: {len(moved)}  added: {len(added)}  "
              f"removed: {len(gone)}  renamed: {len(renamed)}")
        for k in moved:
            print("  ~", k)
        for k in added:
            print("  +", k)
        for k in gone:
            print("  -", k)
        for r in renamed:
            print(f"  → {r['from_path']} → {new_path.get(r['to'], r['to'])}")
        print("\nEvery wiki article citing a changed unit is now stale until re-graded.")
        if renamed:
            print("Renamed units keep their article — repoint citations, do not create a second one.")
        return

    # diff against the previous run BEFORE overwriting the manifest — this is the
    # absorb trigger. Date ranges miss changed code packages, whose entry filename
    # never moves. The raw manifest is the fresher local baseline; a clean clone
    # falls back to the committed twin so the CI queue is scoped to real drift
    # instead of "everything is new".
    prev, prev_path = {}, {}
    base = manifest if manifest.exists() else committed
    if base.exists():
        try:
            _m = json.loads(base.read_text())
            prev = {u["id"]: (u["sha"], u.get("status", "active")) for u in _m}
            prev_path = {u["id"]: u.get("path", "") for u in _m}
        except Exception:
            prev, prev_path = {}, {}
    try:
        from .completion import read_completed
    except ImportError:
        from completion import read_completed
    absorb_log = wiki_dir / "_absorb_log.json"
    absorbed = read_completed(absorb_log)

    pending = build_pending(units, prev, prev_path, absorbed)
    by_id = {u["id"]: u for u in units}

    repo_r = a.repo.resolve()
    edges = load_package_edges(repo_r)
    try:
        head = git(repo_r, "rev-parse", "HEAD").strip()
    except Exception:
        head = None
    if edges:
        print(f"  graph: {len(edges)} packages with cross-package edges")
    a.out.mkdir(parents=True, exist_ok=True)
    written = set()
    for u in units:
        name = f"{(u['first'] or '1970-01-01')[:10]}_{u['id']}.md"
        src = u.get("_inbox_file")
        # feeder entries are already rendered — copy, don't rebuild
        text = (src.read_text(encoding="utf-8", errors="replace") if src is not None
                else render(repo_r, u, edges, head))
        (a.out / name).write_text(text, encoding="utf-8")
        written.add(name)
    for stale in set(os.listdir(a.out)) - written:
        if stale.endswith(".md"):
            os.remove(a.out / stale)

    manifest.parent.mkdir(parents=True, exist_ok=True)
    # drop "files" (bulky) and any _private key (e.g. _inbox_file, a Path)
    payload = json.dumps(
        [{k: v for k, v in u.items() if k != "files" and not k.startswith("_")} for u in units],
        indent=1)
    manifest.write_text(payload)
    committed.parent.mkdir(parents=True, exist_ok=True)
    committed.write_text(payload)

    (manifest.parent / "_pending.json").write_text(json.dumps(pending, indent=1))
    lines = ["# Absorb queue", "",
             "Generated by ingest.py. **This overrides `_absorb_log.json`.**",
             "A `changed` unit is already marked absorbed but its content moved —",
             "re-absorb it and re-grade every article citing it.", ""]
    for bucket, note in (("new", "never absorbed"),
                         ("changed", "sha moved since last ingest — RE-absorb"),
                         ("removed", "gone from the repo — articles citing it are now dead")):
        ids = pending[bucket]
        lines += [f"## {bucket} ({len(ids)}) — {note}", ""]
        if not ids:
            lines += ["_none_", ""]
        for i in ids:
            u = by_id.get(i)
            lines.append(f"- `{i}` — `{u['path']}` ({u['kind']})" if u else f"- `{i}` — _removed_")
        lines.append("")

    ren = pending["renamed"]
    lines += [f"## renamed ({len(ren)}) — same content, new path — REPOINT, do not create", "",
              "The article already exists. Update its `sources:` and every inline",
              "`path@sha` citation to the new path, then bump the sha. Do **not**",
              "write a second article: these are the same documents that appear",
              "nowhere in `new` or `removed`, precisely so they can't be duplicated.", ""]
    if not ren:
        lines += ["_none_", ""]
    for r in ren:
        u = by_id.get(r["to"])
        new_path = u["path"] if u else r["to"]
        lines.append(f"- `{r['from_path']}` → `{new_path}`")
    lines.append("")
    (manifest.parent / "_pending.md").write_text("\n".join(lines))

    kinds = defaultdict(int)
    for u in units:
        kinds[u["kind"]] += 1
    print(f"{len(units)} entries → {a.out}")
    for k, v in sorted(kinds.items(), key=lambda x: -x[1]):
        print(f"  {v:5d}  {k}")
    if wikiignore:
        print(f"  .wikiignore: {len(wikiignore)} pattern(s) — {', '.join(wikiignore)}")
    else:
        print(f"  .wikiignore: none ({repo / WIKIIGNORE_NAME})")
    sup = sum(1 for u in units if u["status"] == "superseded")
    fail = sum(1 for u in units if u["kind"] == "binary_doc")
    print(f"  {sup} marked superseded; {fail} binary docs — check those bodies extracted before absorbing")
    if pending["first_run"]:
        print("\nfirst run — absorb everything")
    else:
        print(f"\nqueue: {len(pending['new'])} new, {len(pending['changed'])} changed, "
              f"{len(pending['removed'])} removed, {len(pending['renamed'])} renamed"
              f"  →  {manifest.parent/'_pending.md'}")
        if pending["changed"]:
            print("  changed units are already in _absorb_log.json. Absorb by queue, not by date.")
        if pending["renamed"]:
            print("  renamed units keep their article — repoint citations, do not create a second one.")


if __name__ == "__main__":
    main()
